"""
Import jadwal teknisi dari Excel ke snc_road_plans (local enterprise DB).
Usage: python3 kil/backend/scripts/import_excel_schedule.py [path_to_xlsx]
"""
import sys
import os
from datetime import datetime

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..', '..'))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..'))

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)

import psycopg
from psycopg.rows import dict_row

XLSX = sys.argv[1] if len(sys.argv) > 1 else "05. JADWAL TEKNISI MEI 2026.xlsx"

# Sheet name -> p_user_id mapping
SHEET_TO_USER = {
    'AKBAR R':   385,   # Akbar Rohmatulah
    'ANAM':      379,   # Choirul Anam
    'ABU S':     396,   # M. Abu Samsudin
    'ALMAS ':    399,   # Ananda Almas
    'ADAM':      403,   # Adam Abdillah
    'MAHRUS':    389,   # Moh. Mahrus
    'ANDIK':     392,   # Andik Noroyan Fananiar
    'LUCKY':     391,   # Lucky Adi Putra
    'MAULANA':   428,   # Moch Maulana
    'RANGGA':    None,  # Not in p_user — skip
    'FATHUR':    None,  # Not in p_user — skip
    'IMAM':      459,   # Nur Imam Siswo Utomo
    'ARGA':      442,   # Argantara Alif Saputra
    'RENDY':     407,   # I Wayan Rendy
    'IRUL':      None,  # Not found — skip
    'MULYASARI': 445,   # Muliyasari
    'PM Jogja ': None,  # PM, skip
    'PM SOLO':   None,  # PM, skip
}

def parse_time(val: str):
    """Parse '13.00-15.00' -> ('13:00', '15:00')"""
    if not val or val in ('-', 'OFF', ''):
        return None, None
    val = val.strip().replace(',', '.')
    if '-' in val:
        parts = val.split('-', 1)
        def fmt(s):
            s = s.strip().replace('.', ':')
            if ':' not in s:
                s = s + ':00'
            return s[:5]
        return fmt(parts[0]), fmt(parts[1])
    return None, None

def parse_sheet(ws):
    """Parse one sheet into list of {date, customer_name, time_start, time_end, service_type}."""
    rows = list(ws.iter_rows(values_only=True))
    visits = []

    i = 0
    while i < len(rows):
        row = rows[i]
        # Detect week header row (contains day names)
        first = str(row[0] or '').strip().upper()
        if first == 'SENIN':
            # Next row has dates
            if i + 1 >= len(rows):
                break
            date_row = rows[i + 1]
            # columns: 0,3,6,9,12 = Mon,Tue,Wed,Thu,Fri
            day_dates = {}
            for col_idx, day in [(0, 0), (3, 1), (6, 2), (9, 3), (12, 4)]:
                cell = date_row[col_idx] if col_idx < len(date_row) else None
                if cell and hasattr(cell, 'date'):
                    day_dates[day] = cell.date()
                elif isinstance(cell, datetime):
                    day_dates[day] = cell.date()
            # Now read job rows until next SENIN or end
            j = i + 2
            while j < len(rows):
                jrow = rows[j]
                jfirst = str(jrow[0] if jrow else '').strip().upper()
                if jfirst == 'SENIN':
                    break
                # Each row has up to 5 days * 3 cols = 15 cols
                for day_idx, col_start in enumerate(range(0, 15, 3)):
                    if day_idx not in day_dates:
                        continue
                    if col_start >= len(jrow):
                        continue
                    raw_cust = jrow[col_start]
                    # Skip cells that contain datetime objects (misplaced dates)
                    if isinstance(raw_cust, datetime):
                        continue
                    cust = str(raw_cust or '').strip()
                    time_val = str(jrow[col_start + 1] if col_start + 1 < len(jrow) else '') or ''
                    svc = str(jrow[col_start + 2] if col_start + 2 < len(jrow) else '') or ''
                    if cust and cust not in ('-', 'OFF', ''):
                        t_start, t_end = parse_time(time_val)
                        visits.append({
                            'date': day_dates[day_idx],
                            'customer_name': cust,
                            'time_start': t_start,
                            'time_end': t_end,
                            'service_type': svc.strip() if svc.strip() not in ('-', '') else None,
                        })
                j += 1
            i = j
        else:
            i += 1

    return visits

def connect_local():
    return psycopg.connect(
        host='127.0.0.1', port=5432,
        dbname='kil_enterprise', user='kil_ent', password='KilEnt2026!',
        row_factory=dict_row,
    )

def main():
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    conn = connect_local()
    cur = conn.cursor()

    # Ensure table exists
    cur.execute("""
        CREATE TABLE IF NOT EXISTS snc_road_plans (
            id               BIGSERIAL PRIMARY KEY,
            p_user_id        INTEGER NOT NULL,
            customer_id      INTEGER,
            customer_name    TEXT,
            visit_date       TIMESTAMPTZ NOT NULL,
            time_start       TEXT,
            time_end         TEXT,
            service_type     TEXT,
            status           TEXT DEFAULT 'Baru',
            notes            TEXT,
            kelava_road_plan_id BIGINT,
            created_at       TIMESTAMPTZ DEFAULT NOW(),
            updated_at       TIMESTAMPTZ DEFAULT NOW()
        )
    """)
    conn.commit()

    total_inserted = 0
    total_skipped = 0

    for sheet_name in wb.sheetnames:
        p_user_id = SHEET_TO_USER.get(sheet_name)
        if p_user_id is None:
            print(f"SKIP sheet '{sheet_name}' (no p_user mapping)")
            continue

        ws = wb[sheet_name]
        visits = parse_sheet(ws)
        print(f"Sheet '{sheet_name}' (p_user={p_user_id}): {len(visits)} visits parsed")

        for v in visits:
            # Check for duplicate
            cur.execute("""
                SELECT id FROM snc_road_plans
                WHERE p_user_id = %s AND visit_date::date = %s AND customer_name = %s
            """, (p_user_id, v['date'], v['customer_name']))
            if cur.fetchone():
                total_skipped += 1
                continue

            visit_dt = datetime.combine(v['date'], datetime.min.time()).replace(hour=8)
            if v['time_start']:
                try:
                    h, m = v['time_start'].split(':')
                    visit_dt = datetime.combine(v['date'], datetime.min.time()).replace(hour=int(h), minute=int(m))
                except Exception:
                    pass

            cur.execute("""
                INSERT INTO snc_road_plans
                    (p_user_id, customer_name, visit_date, time_start, time_end, service_type, status)
                VALUES (%s, %s, %s, %s, %s, %s, 'Baru')
            """, (p_user_id, v['customer_name'], visit_dt, v['time_start'], v['time_end'], v['service_type']))
            total_inserted += 1

        conn.commit()

    print(f"\nDone. Inserted: {total_inserted} | Skipped (duplicates): {total_skipped}")
    conn.close()

if __name__ == '__main__':
    main()
