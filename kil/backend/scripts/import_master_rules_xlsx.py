"""
Import "Data Jadwal Client & Teknisi.xlsx" → snc_recurring_rules.

Sheet 1 "Teknisi": co-visit assignments (Client + tech 1/2/3)
Sheet 2 "Kunjungan Client": recurring rules (Frekuensi, Hari, Jam, Keterangan)

These are MASTER rules dari koordinator (is_mandatory=true).
"""
import argparse
import json
import re
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))

import openpyxl
from kil.db.kelava_db import _get_local_pool
from psycopg.rows import dict_row

DAY_MAP = {
    'senin': 0, 'selasa': 1, 'rabu': 2, 'kamis': 3,
    'jumat': 4, 'sabtu': 5, 'minggu': 6,
}

# Common tech name aliases from xlsx → snc_technicians.name
TECH_ALIASES = {
    'adam': 'Adam Abdillah',
    'akbar': 'Akbar Rohmatulah',
    'almas': 'Ananda Almas',
    'ananda': 'Ananda Almas',
    'anam': 'Choirul Anam',
    'abu': 'M. Abu Samsudin',
    'andik': 'Andik Noroyan Fananiar',
    'lucky': 'Lucky Adi Putra',
    'mahrus': 'Moh. Mahrus',
    'maulana': 'Moch Maulana',
    'arga': 'Argantara Alif Saputra',
    'irul': 'Irul',
    'muliyasari': 'Muliyasari',
    'rendy': 'I Wayan Rendy',
    'fathur': 'Fathur Rozek',
    'imam': 'Nur Imam Siswo Utomo',
    'gilang': 'Muhammad Gilang Isad Abdul Ghofur',
    'khoirul': 'M. Khoirul Anwar',
}


def parse_freq(freq_str: str, note: str) -> tuple[str, str | None]:
    """Returns (frequency, week_pattern)."""
    note = (note or '').lower()
    freq_str = (freq_str or '').lower().strip()

    # Parse "Setiap Minggu ke X" or "Minggu ke X & Y"
    wp = None
    m = re.search(r'minggu ke (\d+)(?:\s*&\s*(\d+))?', note)
    if m:
        nums = [m.group(1)]
        if m.group(2):
            nums.append(m.group(2))
        wp = ','.join(nums)

    # Determine frequency
    if '4x' in freq_str:
        return 'weekly', None  # 4× per month = weekly
    if '2x' in freq_str:
        return 'biweekly', wp or '1,3'
    if '1x' in freq_str:
        return 'monthly', wp or '1'
    return 'weekly', wp


def parse_time(jam: str) -> tuple[str | None, str | None]:
    """'13.00-15.00' → ('13:00', '15:00')."""
    if not jam:
        return None, None
    jam = jam.strip().replace(' ', '').replace('.', ':')
    m = re.match(r'(\d{1,2}):(\d{2})-(\d{1,2}):(\d{2})', jam)
    if not m:
        return None, None
    sh, sm, eh, em = int(m[1]), int(m[2]), int(m[3]), int(m[4])
    sh = sh % 24
    eh = eh % 24
    return f"{sh:02d}:{sm:02d}", f"{eh:02d}:{em:02d}"


def fuzzy_client(name: str, idx: dict, alias_idx: dict) -> int | None:
    cn = re.sub(r'[^a-z0-9]', '', name.lower())
    if cn in idx:
        return idx[cn]
    if cn in alias_idx:
        return alias_idx[cn]
    if len(cn) >= 4:
        for n2, cid in idx.items():
            if len(n2) >= 4 and (cn in n2 or n2 in cn):
                return cid
    return None


def fuzzy_tech(name: str, tech_idx: dict) -> int | None:
    if not name:
        return None
    n = name.strip().lower()
    full = TECH_ALIASES.get(n)
    if full:
        cn = re.sub(r'[^a-z]', '', full.lower())
        return tech_idx.get(cn)
    # Substring match
    cn = re.sub(r'[^a-z]', '', n)
    for tn, tid in tech_idx.items():
        if cn in tn or tn in cn:
            return tid
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx', help='Path to Data Jadwal Client & Teknisi.xlsx')
    ap.add_argument('--dry-run', action='store_true')
    ap.add_argument('--user-id', type=int, default=1)
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)

    # Sheet 1: tech assignments
    sheet_tech = wb['Teknisi']
    co_visits = {}  # client_name → [tech1, tech2, tech3]
    for r in range(4, sheet_tech.max_row + 1):
        name = sheet_tech.cell(r, 2).value
        if not name or not isinstance(name, str):
            continue
        techs = [sheet_tech.cell(r, c).value for c in (3, 4, 5)]
        co_visits[name.strip()] = [t.strip() if t else None for t in techs]

    # Sheet 2: recurring rules
    sheet_rules = wb['Kunjungan Client']
    rules_raw = []
    for r in range(4, sheet_rules.max_row + 1):
        name = sheet_rules.cell(r, 2).value
        freq = sheet_rules.cell(r, 3).value
        hari = sheet_rules.cell(r, 4).value
        jam = sheet_rules.cell(r, 5).value
        ket = sheet_rules.cell(r, 6).value
        if not name or not isinstance(name, str):
            continue
        rules_raw.append({
            'client_name': name.strip(),
            'freq': freq, 'hari': hari, 'jam': jam, 'ket': ket,
        })

    print(f'Found {len(co_visits)} co-visit assignments + {len(rules_raw)} rule rows')

    with _get_local_pool().connection() as conn:
        with conn.cursor(row_factory=dict_row) as cur:
            # Load clients + aliases
            cur.execute("SELECT id, name FROM snc_clients")
            cli_rows = cur.fetchall()
            idx_client = {re.sub(r'[^a-z0-9]', '', r['name'].lower()): r['id'] for r in cli_rows}
            cur.execute("SELECT alias, client_id FROM snc_client_aliases")
            alias_idx = {re.sub(r'[^a-z0-9]', '', r['alias'].lower()): r['client_id'] for r in cur.fetchall()}

            # Load techs
            cur.execute("SELECT id, name FROM snc_technicians WHERE is_active = true")
            tech_idx = {re.sub(r'[^a-z]', '', r['name'].lower()): r['id'] for r in cur.fetchall()}

            stats = {'rules_inserted': 0, 'rules_updated': 0,
                     'unmatched_client': [], 'unmatched_tech': [],
                     'co_visit_applied': 0}

            for rr in rules_raw:
                client_id = fuzzy_client(rr['client_name'], idx_client, alias_idx)
                if not client_id:
                    stats['unmatched_client'].append(rr['client_name'])
                    if args.dry_run:
                        print(f"  [DRY] UNMATCHED client: {rr['client_name']}")
                    continue

                day_idx = DAY_MAP.get((rr['hari'] or '').lower().strip())
                if day_idx is None:
                    print(f"  WARN bad hari for {rr['client_name']}: {rr['hari']}")
                    continue

                freq, wp = parse_freq(rr['freq'], rr['ket'])
                t_start, t_end = parse_time(rr['jam'])

                # Pick primary + backups from co_visit sheet
                co = co_visits.get(rr['client_name'])
                if not co:
                    # Try fuzzy match on co-visit
                    for cv_name, cv_techs in co_visits.items():
                        cn1 = re.sub(r'[^a-z0-9]', '', rr['client_name'].lower())
                        cn2 = re.sub(r'[^a-z0-9]', '', cv_name.lower())
                        if cn1 == cn2 or cn1 in cn2 or cn2 in cn1:
                            co = cv_techs
                            break
                if co:
                    primary_id = fuzzy_tech(co[0], tech_idx) if co[0] else None
                    bkp1_id = fuzzy_tech(co[1], tech_idx) if co[1] else None
                    bkp2_id = fuzzy_tech(co[2], tech_idx) if co[2] else None
                    stats['co_visit_applied'] += 1
                else:
                    primary_id, bkp1_id, bkp2_id = None, None, None

                if not primary_id:
                    print(f"  WARN no primary tech for {rr['client_name']} — looking up history")
                    # Fallback: look up tech from past visits
                    cur.execute("""
                        SELECT technician_id, COUNT(*) AS n
                        FROM snc_schedule_events
                        WHERE client_id = %s AND start_date >= '2026-03-01'
                        GROUP BY technician_id ORDER BY n DESC LIMIT 1
                    """, (client_id,))
                    row = cur.fetchone()
                    if row:
                        primary_id = row['technician_id']
                    else:
                        print(f"    SKIP no history either: {rr['client_name']}")
                        continue

                if args.dry_run:
                    print(f"  [DRY] {rr['client_name']:25} day={day_idx} time={t_start}-{t_end} "
                          f"freq={freq} wp={wp} primary={primary_id} bkp1={bkp1_id} bkp2={bkp2_id}")
                    stats['rules_inserted'] += 1
                    continue

                # UPSERT rule
                notes = f"Master rule dari xlsx koord. Freq raw={rr['freq']} ket={rr['ket'] or ''}"
                try:
                    cur.execute("""
                        INSERT INTO snc_recurring_rules
                            (client_id, primary_tech_id, backup_tech_1_id, backup_tech_2_id,
                             frequency, weekdays, week_pattern,
                             time_start, time_end, visit_type,
                             is_mandatory, suppress_holiday,
                             notes, effective_start, created_by)
                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, 'PRC',
                                true, true, %s, %s, %s)
                        ON CONFLICT (client_id, effective_start) DO UPDATE SET
                            primary_tech_id = EXCLUDED.primary_tech_id,
                            backup_tech_1_id = EXCLUDED.backup_tech_1_id,
                            backup_tech_2_id = EXCLUDED.backup_tech_2_id,
                            frequency = EXCLUDED.frequency,
                            weekdays = EXCLUDED.weekdays,
                            week_pattern = EXCLUDED.week_pattern,
                            time_start = EXCLUDED.time_start,
                            time_end = EXCLUDED.time_end,
                            is_mandatory = true,
                            notes = EXCLUDED.notes,
                            updated_at = now()
                        RETURNING id, xmax = 0 AS is_new
                    """, (client_id, primary_id, bkp1_id, bkp2_id,
                          freq, [day_idx], wp,
                          t_start, t_end,
                          notes, date.today().isoformat(), args.user_id))
                    res = cur.fetchone()
                    if res['is_new']:
                        stats['rules_inserted'] += 1
                    else:
                        stats['rules_updated'] += 1
                except Exception as e:
                    print(f"  ERR rule {rr['client_name']}: {e}")

            if args.dry_run:
                conn.rollback()
            else:
                conn.commit()

            print(f"\nResults: {stats['rules_inserted']} new, {stats['rules_updated']} updated, "
                  f"{stats['co_visit_applied']} with co-visit, "
                  f"{len(stats['unmatched_client'])} unmatched client")
            if stats['unmatched_client']:
                print('Unmatched:', stats['unmatched_client'])


if __name__ == '__main__':
    main()
