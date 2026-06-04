"""
Audit: enumerate semua rule di master xlsx, lalu verify generate-draft
hasilkan jumlah visit yang BENAR — account holiday + co-visit + week_pattern.

Usage:
  audit_multislot_rules.py <xlsx_path>                    # check rule existence
  audit_multislot_rules.py <xlsx_path> --month 2026-06    # check visit counts
"""
import argparse
import re
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))

import openpyxl
from kil.db.kelava_db import _get_local_pool
from psycopg.rows import dict_row

DAY = {'senin': 0, 'selasa': 1, 'rabu': 2, 'kamis': 3,
       'jumat': 4, 'sabtu': 5, 'minggu': 6}


def expected_dates_for_rule(year: int, month: int, dow: int,
                            freq: str, ket: str, suppressed: set) -> list:
    """Compute expected target dates for one xlsx rule row."""
    import calendar as cal
    _, ndays = cal.monthrange(year, month)
    freq_l = (freq or '').lower()
    ket_l = (ket or '').lower()

    wp = None
    m = re.search(r'minggu ke (\d+)(?:\s*&\s*(\d+))?', ket_l)
    if m:
        wp = [int(m[1])] + ([int(m[2])] if m[2] else [])

    if '4x' in freq_l:
        target_woms = list(range(1, 6))   # weekly = semua minggu
    elif '2x' in freq_l:
        target_woms = wp or [1, 3]
    elif '1x' in freq_l:
        target_woms = wp or [1]
    else:
        target_woms = list(range(1, 6))

    out = []
    for day in range(1, ndays + 1):
        d = date(year, month, day)
        if d.weekday() != dow:
            continue
        wom = (d.day - 1) // 7 + 1
        if wom not in target_woms:
            continue
        if d in suppressed:
            continue
        out.append(d)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx')
    ap.add_argument('--month', help='YYYY-MM untuk audit visit counts')
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    sh = wb['Kunjungan Client']

    all_rules = []
    by_client = defaultdict(list)
    for r in range(4, sh.max_row + 1):
        name = sh.cell(r, 2).value
        if not name or not isinstance(name, str):
            continue
        row = {
            'name': name.strip(),
            'freq': (sh.cell(r, 3).value or '').strip(),
            'hari': (sh.cell(r, 4).value or '').strip(),
            'jam':  (sh.cell(r, 5).value or '').strip(),
            'ket':  (sh.cell(r, 6).value or '').strip(),
        }
        all_rules.append(row)
        by_client[name.strip()].append(row)

    multi = {n: rows for n, rows in by_client.items() if len(rows) > 1}
    print(f'Multi-slot clients: {len(multi)} dari {len(by_client)} total')

    with _get_local_pool().connection() as conn:
        with conn.cursor(row_factory=dict_row) as c:
            if not args.month:
                print('\n' + '═' * 70)
                print('Verifikasi DB — apakah semua slot ke-import?')
                print('═' * 70)
                all_ok = True
                for n, rows in sorted(multi.items()):
                    c.execute("""SELECT COUNT(*) AS n FROM snc_recurring_rules r
                                 JOIN snc_clients c ON c.id = r.client_id
                                 WHERE c.name ILIKE %s AND r.is_mandatory = true
                                   AND (r.effective_end IS NULL OR r.effective_end >= CURRENT_DATE)""",
                              (f'%{n.split()[0]}%',))
                    got = c.fetchone()['n']
                    status = '✓' if got >= len(rows) else '✗'
                    if got < len(rows):
                        all_ok = False
                    print(f'  {status} {n:25s} xlsx={len(rows)} db={got}')
                print('\nALL OK' if all_ok else '\n⚠️  SOME MISSING — re-run import')
                return

            # --month mode — verify actual visit counts
            year, mo = map(int, args.month.split('-'))
            c.execute("""SELECT id FROM snc_draft_batches WHERE target_month=%s
                         ORDER BY id DESC LIMIT 1""", (args.month,))
            br = c.fetchone()
            if not br:
                print(f'No draft batch untuk {args.month}'); return
            bid = br['id']

            c.execute("""SELECT suppression_date FROM snc_suppression_dates
                         WHERE EXTRACT(YEAR FROM suppression_date)=%s
                           AND EXTRACT(MONTH FROM suppression_date)=%s""", (year, mo))
            suppressed = {r['suppression_date'] for r in c.fetchall()}
            print(f'\nHari libur {args.month}: {sorted(suppressed)}')

            print('\n' + '═' * 70)
            print(f'Verifikasi expected vs actual visit counts (batch {bid})')
            print('═' * 70)
            perfect = under = over = 0
            for x in all_rules:
                dow = DAY.get(x['hari'].lower())
                if dow is None:
                    continue
                tm = re.match(r'(\d{1,2})\.(\d{2})', x['jam'].split('-')[0].strip())
                if not tm:
                    continue
                h, mm = int(tm[1]) % 24, int(tm[2])
                exp_dates = expected_dates_for_rule(year, mo, dow, x['freq'],
                                                     x['ket'], suppressed)

                key = re.sub(r'[^a-z0-9]', '', x['name'].lower())[:5]
                c.execute("""SELECT primary_tech_id, backup_tech_1_id, backup_tech_2_id
                             FROM snc_recurring_rules r JOIN snc_clients c ON c.id=r.client_id
                             WHERE LOWER(REGEXP_REPLACE(c.name,'[^a-z0-9]','','gi')) LIKE %s
                               AND r.is_mandatory=true LIMIT 1""", (f'%{key}%',))
                tr = c.fetchone()
                ntech = max(sum(1 for t in (tr['primary_tech_id'],
                                            tr['backup_tech_1_id'],
                                            tr['backup_tech_2_id']) if t), 1) if tr else 1
                exp = len(exp_dates) * ntech

                c.execute("""SELECT COUNT(*) AS n FROM snc_schedule_events se
                             JOIN snc_clients c ON c.id=se.client_id
                             WHERE LOWER(REGEXP_REPLACE(c.name,'[^a-z0-9]','','gi')) LIKE %s
                               AND se.draft_batch_id=%s AND EXTRACT(DOW FROM se.start_date)=%s
                               AND EXTRACT(HOUR FROM se.start_datetime)=%s
                               AND EXTRACT(MINUTE FROM se.start_datetime)=%s""",
                          (f'%{key}%', bid, (dow + 1) % 7, h, mm))
                got = c.fetchone()['n']

                if got == exp:
                    perfect += 1
                elif got > exp:
                    over += 1
                    print(f'  over  {x["name"]:18s} {x["hari"]:7s} {x["jam"]:14s} exp={exp} got={got}')
                else:
                    under += 1
                    print(f'  under {x["name"]:18s} {x["hari"]:7s} {x["jam"]:14s} exp={exp} got={got}')
            print(f'\n→ {perfect}/{len(all_rules)} exact match | {under} under | {over} over')


if __name__ == '__main__':
    main()
