"""
Import Omzet Invoice file → snc_contracts (SSOT for active customers).

Per user rule:
  - Hanya customer di Omzet file yang jadi SSOT
  - Customer hanya boleh di-schedule untuk bulan X jika X ada di range
    first_invoice_month .. last_invoice_month customer tersebut.

Process:
  1. Parse xlsx → per customer: months with invoice + total
  2. Fuzzy-match ke snc_clients; create new client kalau belum ada
  3. UPSERT snc_contracts (one per customer) dengan:
       start_date = first day of first invoiced month
       end_date   = last day of last invoiced month
       nilai_kontrak = sum total
       no_kontrak = 'OMZET-2026-<client_id>'
  4. Flag snc_clients yang tidak ada di Omzet → is_active=false (optional)

Usage:
    .venv/bin/python -m kil.backend.scripts.import_omzet_ssot \
        Omzet_Invoice_2026_1780371702.xlsx [--year 2026] [--dry-run]
"""
import argparse
import calendar
import json
import re
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))

import openpyxl
from kil.db.kelava_db import _get_local_pool
from psycopg.rows import dict_row

MONTH_COLS = {4:1, 5:2, 6:3, 7:4, 8:5, 9:6, 10:7, 11:8, 12:9, 13:10, 14:11, 15:12}


def parse_omzet(xlsx_path: str) -> dict:
    """Returns {name: {months: set, total: float}}"""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    customers = defaultdict(lambda: {"months": set(), "total": 0.0})
    for r in range(4, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if not name or not isinstance(name, str): continue
        if name.strip().lower() in ('name', 'subitems'): continue
        name = name.strip()
        for col, m in MONTH_COLS.items():
            v = ws.cell(r, col).value
            if isinstance(v, (int, float)) and v > 0:
                customers[name]["months"].add(m)
                customers[name]["total"] += v
    # Drop entries with no invoice
    return {n: d for n, d in customers.items() if d["months"]}


def fuzzy_match(name: str, idx: dict) -> int | None:
    cn = re.sub(r'[^a-z0-9]', '', name.lower())
    if cn in idx:
        return idx[cn]
    if len(cn) >= 5:
        for n2, cid in idx.items():
            if len(n2) >= 5 and (cn in n2 or n2 in cn):
                return cid
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--year", type=int, default=2026)
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--deactivate-orphan-clients", action="store_true",
                    help="Set is_active=false untuk snc_clients yang tidak ada di Omzet")
    args = ap.parse_args()

    print(f"Parsing {args.xlsx}…")
    omzet = parse_omzet(args.xlsx)
    print(f"  {len(omzet)} customers ditemukan")

    with _get_local_pool().connection() as conn:
        with conn.cursor(row_factory=dict_row) as cur:
            cur.execute("SELECT id, name FROM snc_clients")
            existing = cur.fetchall()
            idx = {re.sub(r'[^a-z0-9]', '', r['name'].lower()): r['id']
                   for r in existing}

            stats = {"matched_exact": 0, "matched_fuzzy": 0, "new_clients": 0,
                     "contracts_upserted": 0, "skipped": 0}
            omzet_client_ids = set()

            for name, d in omzet.items():
                # Match or create
                cid = fuzzy_match(name, idx)
                if cid:
                    cn = re.sub(r'[^a-z0-9]', '', name.lower())
                    if cn in idx:
                        stats["matched_exact"] += 1
                    else:
                        stats["matched_fuzzy"] += 1
                else:
                    if args.dry_run:
                        stats["new_clients"] += 1
                        print(f"  [DRY] CREATE client: {name}")
                        continue
                    # Create new client
                    try:
                        cur.execute("""
                            INSERT INTO snc_clients (name, is_active, notes)
                            VALUES (%s, true, %s)
                            ON CONFLICT (name) DO NOTHING
                            RETURNING id
                        """, (name, "Auto-created dari Omzet Invoice 2026"))
                        row = cur.fetchone()
                        if row:
                            cid = row['id']
                            stats["new_clients"] += 1
                            idx[re.sub(r'[^a-z0-9]', '', name.lower())] = cid
                        else:
                            cur.execute("SELECT id FROM snc_clients WHERE name = %s", (name,))
                            cid = cur.fetchone()['id']
                            stats["matched_exact"] += 1
                    except Exception as e:
                        print(f"  ERR creating client {name}: {e}")
                        stats["skipped"] += 1
                        continue

                omzet_client_ids.add(cid)

                # Compute contract period
                months = sorted(d["months"])
                first_m, last_m = months[0], months[-1]
                start_date = date(args.year, first_m, 1)
                _, ndays = calendar.monthrange(args.year, last_m)
                end_date = date(args.year, last_m, ndays)
                no_kontrak = f"OMZET-{args.year}-{cid}"

                if args.dry_run:
                    stats["contracts_upserted"] += 1
                    continue

                # UPSERT contract
                try:
                    cur.execute("""
                        INSERT INTO snc_contracts
                            (snc_customer_id, no_kontrak, start_date, end_date,
                             is_active, nilai_kontrak, frekuensi_visit, notes, created_by)
                        VALUES (%s, %s, %s, %s, 'YES', %s, 1, %s, 1)
                        ON CONFLICT DO NOTHING
                        RETURNING id
                    """, (cid, no_kontrak, start_date, end_date, d["total"],
                          f"Auto-import dari Omzet Invoice {args.year}. "
                          f"Active months: {months}"))
                    row = cur.fetchone()
                    if row:
                        stats["contracts_upserted"] += 1
                    else:
                        # Update existing
                        cur.execute("""
                            UPDATE snc_contracts
                            SET start_date = LEAST(start_date, %s),
                                end_date   = GREATEST(end_date, %s),
                                nilai_kontrak = %s,
                                notes = %s,
                                updated_at = now()
                            WHERE no_kontrak = %s
                        """, (start_date, end_date, d["total"],
                              f"Updated from Omzet {args.year}. Months: {months}",
                              no_kontrak))
                        stats["contracts_upserted"] += 1
                except Exception as e:
                    print(f"  ERR upserting contract {name}: {e}")

            if args.dry_run:
                conn.rollback()
                print(f"\n[DRY RUN] {stats}")
            else:
                conn.commit()
                print(f"\n✓ {stats}")

            # Optional: deactivate orphan clients
            if args.deactivate_orphan_clients and not args.dry_run:
                orphan_ids = [r['id'] for r in existing
                              if r['id'] not in omzet_client_ids]
                # Only deactivate if currently active
                cur.execute("""
                    UPDATE snc_clients
                    SET is_active = false,
                        deactivated_at = now(),
                        deactivation_reason = 'Tidak ada di Omzet Invoice 2026 (SSOT)'
                    WHERE id = ANY(%s) AND COALESCE(is_active, true) = true
                """, (orphan_ids,))
                conn.commit()
                print(f"  Orphan clients deactivated: {cur.rowcount}")


if __name__ == "__main__":
    main()
